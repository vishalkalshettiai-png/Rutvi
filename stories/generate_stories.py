#!/usr/bin/env python3
"""Generate 100 motivational stories for Class 2 students as Word and other formats."""

from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

STORIES = [
    (
        "The Little Seed That Waited",
        "A tiny seed sat under the soil. It was dark and cold. The seed wanted to grow fast, but it waited patiently. Rain came, sun shone, and slowly a green sprout appeared. Soon it became a strong plant. Moral: Be patient. Good things take time.",
    ),
    (
        "Ravi and the Broken Pencil",
        "Ravi's pencil broke while he was writing. He felt sad and wanted to stop. His friend said, \"Try again with a new pencil.\" Ravi sharpened another pencil and finished his work. His teacher smiled. Moral: Don't give up when something breaks. Try again.",
    ),
    (
        "The Ant and the Big Crumb",
        "A small ant found a big crumb. It was heavy, but the ant did not leave it. Step by step, the ant carried the crumb home. Other ants clapped with their tiny feet. Moral: Hard work can move big things.",
    ),
    (
        "Meena Learns to Tie Her Laces",
        "Meena could not tie her shoelaces. She tried every morning and failed. One day she practised slowly with her mother. Loop, pull, and done! Meena danced with joy. Moral: Practice makes you better.",
    ),
    (
        "The Brave Sparrow",
        "A baby sparrow was afraid to fly. Mother sparrow said, \"I am with you.\" The baby flapped once, then twice, and flew to the next branch. Soon it flew across the garden. Moral: Be brave. You can do more than you think.",
    ),
    (
        "Sharing the Last Biscuit",
        "Aarav had one biscuit left. His sister looked hungry. Aarav broke the biscuit into two and shared it. Both children smiled. Moral: Sharing makes happiness grow.",
    ),
    (
        "The Polite Parrot",
        "A parrot learned two new words: \"Please\" and \"Thank you.\" Whenever someone gave it food, it said thank you. Everyone loved the polite parrot. Moral: Kind words make friends.",
    ),
    (
        "Kavya and the Spelling Test",
        "Kavya missed three spellings in a test. She felt upset. That evening she wrote each word ten times. Next week she got full marks. Moral: Mistakes help us learn.",
    ),
    (
        "The Clean Classroom",
        "Papers were lying on the floor. Rohan picked them up even though he did not drop them. His classmates joined him. The room looked bright again. Moral: Helping without being asked is real kindness.",
    ),
    (
        "Little Lamp in the Dark",
        "The lights went off at night. A small lamp glowed in the corner. It was not very bright, but it guided everyone safely. Moral: Even a small light can help a lot.",
    ),
    (
        "Sana Tries Again",
        "Sana fell while riding her bicycle. She cried a little, then stood up. She tried again and again. Soon she rode without falling. Moral: Falling is okay. Getting up is what matters.",
    ),
    (
        "The Honest Shopkeeper's Son",
        "Vikram found an extra coin at the shop. He gave it back to the customer. The customer patted his head and said, \"You are honest.\" Moral: Honesty is a treasure.",
    ),
    (
        "Two Friends and One Umbrella",
        "Rain started suddenly. Priya had an umbrella, but Anil did not. Priya shared her umbrella and both stayed dry. Moral: Friends share and care.",
    ),
    (
        "The Slow Tortoise Wins Again",
        "A tortoise and a rabbit raced. The rabbit ran fast, then slept. The tortoise walked slowly without stopping. The tortoise reached first. Moral: Slow and steady wins the race.",
    ),
    (
        "Planting a Tree",
        "On her birthday, Diya planted a small tree. She watered it every day. Years later it gave cool shade to many people. Moral: Good deeds grow like trees.",
    ),
    (
        "The Boy Who Said Sorry",
        "Kabir pushed his friend by mistake. He said, \"I am sorry,\" and helped him up. They played together again. Moral: A true sorry can fix a hurt.",
    ),
    (
        "Listening Ears",
        "Teacher told a story. Some children talked, but Neha listened carefully. Later she answered every question. Moral: Listening helps you learn.",
    ),
    (
        "The Patient Kitten",
        "A kitten wanted milk, but Mother cat said, \"Wait until it cools.\" The kitten waited quietly. Soon the milk was safe and tasty. Moral: Patience brings good results.",
    ),
    (
        "Colouring Inside the Lines",
        "Isha's colouring went outside the lines at first. She slowed down and tried again. Her picture looked neat and pretty. Moral: Careful work looks beautiful.",
    ),
    (
        "Morning Smile",
        "Arjun woke up late and felt grumpy. He washed his face, smiled at the mirror, and said, \"I can have a good day.\" At school he made two new friends. Moral: A smile can change your day.",
    ),
    (
        "The Magical Magic Word",
        "When Tina said \"Please,\" doors seemed to open. People helped her faster. She learned that polite words are like magic. Moral: Please and thank you are powerful.",
    ),
    (
        "Carrying Water Together",
        "Two brothers had to fill a big pot. Alone it was hard. Together they carried it easily. Moral: Teamwork makes work light.",
    ),
    (
        "The Lost Eraser",
        "Mira lost her eraser and cried. Then she looked under the desk, inside the bag, and near the window. She found it! Moral: Look carefully before you worry.",
    ),
    (
        "A Cup of Water for Grandpa",
        "Grandpa was tired. Little Sam brought him a cup of water without being told. Grandpa hugged him tightly. Moral: Caring for elders is love.",
    ),
    (
        "The Star That Kept Shining",
        "Clouds covered the sky. One tiny star still peeked through. A child saw it and felt hope. Moral: Keep shining even when things look dark.",
    ),
    (
        "Reading One Page a Day",
        "Yash did not like reading. He started with just one page a day. After a month he finished a whole storybook. Moral: Small steps make big progress.",
    ),
    (
        "The Girl Who Returned the Book",
        "Ananya borrowed a library book and returned it on time. The librarian trusted her with more books. Moral: Being responsible earns trust.",
    ),
    (
        "Helping the New Student",
        "A new boy sat alone. Riya showed him the classroom and shared her lunch. Soon he felt at home. Moral: Welcome others with kindness.",
    ),
    (
        "The Balloon That Came Back",
        "A balloon floated away. The child did not shout. He waited calmly and it got stuck in a low tree. Father brought it down. Moral: Stay calm when things go wrong.",
    ),
    (
        "Saving a Butterfly",
        "A butterfly was trapped near a window. Softly, Zoya opened the window and set it free. It flew into the sunshine. Moral: Be gentle with living things.",
    ),
    (
        "The First Day Fear",
        "On the first day of school, Aman was scared. His teacher held his hand and smiled. By evening Aman said, \"School is fun!\" Moral: New places become friendly with time.",
    ),
    (
        "Cleaning After Play",
        "After playing with blocks, Sara put each block back in the box. Mother was happy and gave her a hug. Moral: Finish your work by cleaning up.",
    ),
    (
        "The Loud Drum and Soft Song",
        "A boy loved beating his drum loudly. One day he learned a soft song on a flute. People listened quietly and smiled. Moral: Soft and calm can be strong too.",
    ),
    (
        "Waiting for Your Turn",
        "At the slide, children pushed. Only Leela waited in line. When her turn came, she slid happily and safely. Moral: Waiting your turn is fair.",
    ),
    (
        "The Broken Toy Fixed",
        "A toy car broke. Instead of crying all day, Dev and Father glued it carefully. The car rolled again. Moral: Fix what you can with care.",
    ),
    (
        "Saying No to Cheating",
        "During a test, a friend whispered answers. Ishaan said softly, \"No, I will try myself.\" He scored less but felt proud. Moral: Honest work is real success.",
    ),
    (
        "The Rainy Day Plan",
        "Rain stopped outdoor play. The children made paper boats and told stories indoors. The day became special. Moral: Make the best of what you have.",
    ),
    (
        "Grandmother's Story Time",
        "Every night Grandmother told a story. One night the child told a story back. Grandmother clapped. Moral: Learning and sharing go together.",
    ),
    (
        "The Thirsty Crow's Idea",
        "A crow found a pot with little water. It dropped pebbles in until the water rose. Then it drank happily. Moral: Think smart when you have a problem.",
    ),
    (
        "Brushing Without Reminders",
        "Earlier Mother reminded Nisha to brush. One morning Nisha brushed before anyone asked. Mother was proud. Moral: Good habits start with you.",
    ),
    (
        "The Tall Ladder of Goals",
        "A kitten wanted to climb a tall shelf. It climbed one step at a time on a small stool, then another. At last it reached. Moral: Reach big goals step by step.",
    ),
    (
        "Quiet in the Library",
        "Friends wanted to talk loudly. Farhan put a finger on his lips and whispered. They read peacefully. Moral: Respect quiet places.",
    ),
    (
        "The Extra Chapati",
        "Mother packed an extra chapati. At school, a classmate had forgotten lunch. The child shared the extra food. Moral: Think of others.",
    ),
    (
        "Learning to Swim",
        "Water scared Pooja. With a coach she practised kicking and floating. One day she swam across the pool. Moral: Face your fears with help and practice.",
    ),
    (
        "The Wobbly Table",
        "A table wobbled in class. Two students put a folded paper under the leg. It stood steady. Moral: Small fixes can solve big wobbles.",
    ),
    (
        "Thank You Note",
        "After a birthday gift, Om wrote a thank-you note to his uncle. Uncle felt loved and sent a happy reply. Moral: Gratitude makes bonds stronger.",
    ),
    (
        "The Early Bird",
        "Reena woke up early, packed her bag, and reached school on time. She felt fresh all day. Moral: Being early reduces worry.",
    ),
    (
        "Painting with Friends",
        "Three friends painted one big mural. Each painted a part. Together the picture looked complete. Moral: Different talents make one beautiful whole.",
    ),
    (
        "The Lost Puppy",
        "A puppy was lost near the park. Children asked adults for help and found its owner. Moral: Ask for help when you need it.",
    ),
    (
        "No Fighting Over the Ball",
        "Two boys fought for one ball. A girl suggested they take turns. They played longer and laughed more. Moral: Taking turns ends fights.",
    ),
    (
        "The Sticky Glue Lesson",
        "Harsh used too much glue and made a mess. Next time he used only a little. His craft was neat. Moral: Learn from messy mistakes.",
    ),
    (
        "Standing Up for a Friend",
        "Someone teased Maya. Her friend said firmly, \"Please stop. That is not kind.\" The teasing stopped. Moral: Defend friends with brave words.",
    ),
    (
        "The Garden Helper",
        "Weeds grew in the school garden. Class 2 pulled weeds and watered plants. Flowers bloomed again. Moral: Care turns places beautiful.",
    ),
    (
        "Counting Coins for Charity",
        "Children saved small coins in a box. They donated them to help needy kids get books. Moral: Little savings can help others.",
    ),
    (
        "The Sleepy Owl Learns Rest",
        "A young owl stayed awake too long and felt weak. Mother owl taught it to rest well. Next night it flew strongly. Moral: Rest helps you grow strong.",
    ),
    (
        "Finishing Homework First",
        "TV called loudly, but Tara finished homework first. Then she watched her show without worry. Moral: Work first, play next.",
    ),
    (
        "The Magical Manners Mirror",
        "A boy frowned at the mirror and it looked gloomy. He smiled and said hello. The day felt brighter. Moral: Your manners shape your world.",
    ),
    (
        "Crossing the Road Safely",
        "Cars zoomed past. A child waited for the green signal, held an adult's hand, and crossed safely. Moral: Safety rules protect you.",
    ),
    (
        "The Broken Promise Fixed",
        "Ritu promised to return a crayon but forgot. Next day she returned it and said sorry. Friendship stayed strong. Moral: Keep promises, and fix them if you forget.",
    ),
    (
        "Singing Softly for Baby",
        "The baby was crying. Big sister sang a soft lullaby. The baby slept. Moral: Gentle actions can calm others.",
    ),
    (
        "The Paper Plane Champion",
        "Many paper planes fell quickly. One boy folded carefully and tested again. His plane flew farthest. Moral: Careful practice beats rushed tries.",
    ),
    (
        "Saying Hello to Neighbours",
        "A shy girl practised saying \"Hello\" to neighbours. Soon everyone knew her name and waved back. Moral: Friendly hellos open doors.",
    ),
    (
        "The Empty Water Bottle",
        "During sports day a child felt thirsty. A classmate shared water from a bottle. Both felt happier. Moral: Share what you can.",
    ),
    (
        "Learning Tables",
        "Multiplication tables seemed hard. Hari chanted them while walking to school. In a week he knew them well. Moral: Repeat to remember.",
    ),
    (
        "The Nest Builders",
        "Two birds built a nest together. One brought twigs, the other soft grass. Their home became cozy. Moral: Help each other build good things.",
    ),
    (
        "Not Laughing at Mistakes",
        "A classmate misread a word. Instead of laughing, the class waited kindly. The child tried again and got it right. Moral: Kindness helps courage grow.",
    ),
    (
        "The Raincoat Ready",
        "Sky looked cloudy. Mother packed a raincoat. When rain came, the child stayed dry and glad. Moral: Being prepared is smart.",
    ),
    (
        "Drawing Every Day",
        "A girl wanted to draw like an artist. She drew a little every day. After months her sketches looked wonderful. Moral: Daily practice creates talent.",
    ),
    (
        "The Quiet Helper",
        "Without telling anyone, a boy sharpened pencils for the class. Teacher noticed and thanked him. Moral: Quiet help is still great help.",
    ),
    (
        "Respecting Nature",
        "Children saw flowers and wanted to pluck them all. Teacher said, \"Leave some for bees and beauty.\" They plucked only one each. Moral: Take only what you need from nature.",
    ),
    (
        "The Sticky Situation",
        "Jam spilled on the table. Instead of hiding it, the child cleaned it with a cloth. Mother appreciated the honesty. Moral: Own your mess and clean it.",
    ),
    (
        "Winning and Losing Gracefully",
        "In a race, one child came first and one came last. The winner congratulated the loser for trying. Both smiled. Moral: Be graceful in win and loss.",
    ),
    (
        "The Birthday Without Gifts",
        "A poor friend had no gift to bring. He sang a song instead. Everyone clapped louder than for toys. Moral: Love is the best gift.",
    ),
    (
        "Saving Electricity",
        "Lights were on in an empty room. A child switched them off. Father said the house felt proud. Moral: Saving power is caring for Earth.",
    ),
    (
        "The Courage to Ask",
        "A sum was confusing. The shy student raised a hand and asked. Teacher explained, and many others understood too. Moral: Asking questions helps everyone.",
    ),
    (
        "Walking the Dog Daily",
        "A puppy needed exercise. The child walked it every evening, rain or shine. The puppy stayed healthy and happy. Moral: Responsibility means doing it every day.",
    ),
    (
        "The Two Boots",
        "Left boot and right boot argued who was more important. Without both, the child could not walk well. They worked as a pair. Moral: We need each other.",
    ),
    (
        "Making the Bed",
        "Every morning after waking, a child made the bed neatly. The room looked ready for a good day. Moral: Small morning habits set a strong start.",
    ),
    (
        "The Forgiving Heart",
        "A friend broke a crayon by accident. Instead of staying angry, the child forgave and shared another. Play continued. Moral: Forgiveness keeps friendship alive.",
    ),
    (
        "Learning a New Dance Step",
        "The dance step looked hard. The child practised in front of a mirror after school. On stage the step was perfect. Moral: Rehearse until you feel ready.",
    ),
    (
        "The Useful Old Bottle",
        "An empty bottle was almost thrown away. The class turned it into a flower vase. Moral: Old things can find new uses.",
    ),
    (
        "Speaking Truth Gently",
        "A vase was cracked. The child told Mother the truth softly and offered to help fix it. Mother hugged the child. Moral: Tell the truth with kindness.",
    ),
    (
        "The Marathon of Reading",
        "A class read for fifteen minutes daily for a month. Together they finished many books. Moral: Together, steady habits win.",
    ),
    (
        "Helping in the Kitchen",
        "While Mother cooked, the child washed vegetables carefully. Dinner was ready faster. Moral: Helping at home shows love.",
    ),
    (
        "The Cloud That Shared Rain",
        "A cloud floated over dry fields and shared its rain. Crops grew green again. Moral: Sharing what you have brings life.",
    ),
    (
        "Not Comparing Scores",
        "Two friends got different marks. They cheered each other instead of comparing. Both studied better next time. Moral: Compete with yourself, support your friends.",
    ),
    (
        "The Sticky Note Reminder",
        "A forgetful child put a sticky note on the bag: \"Bring notebook.\" It worked! Moral: Simple reminders help good habits.",
    ),
    (
        "Building a Sandcastle Again",
        "Waves washed the sandcastle away. The children built a stronger one farther from the water. Moral: Rebuild smarter after setbacks.",
    ),
    (
        "The Whisper of Encouragement",
        "Before a stage play, a nervous child heard a friend whisper, \"You can do it.\" The child performed well. Moral: Encouraging words give courage.",
    ),
    (
        "Keeping Secrets Kindly",
        "A friend shared a worry. The child listened and did not gossip. Trust grew. Moral: Keep private talks private.",
    ),
    (
        "The Ladder of Kind Acts",
        "One kind act a day: smile, share, help, thank. In a week the child felt happier. Moral: Kindness is a daily practice.",
    ),
    (
        "Learning from Younger Sibling",
        "A big brother taught drawing, but the little sister taught him how to laugh freely. Both learned. Moral: Everyone can teach something.",
    ),
    (
        "The Umbrella of Hope",
        "On a stormy day an old woman had no cover. A child offered to walk her home under an umbrella. Moral: Hope walks with helpful hands.",
    ),
    (
        "Finishing What You Start",
        "A puzzle had many pieces. The child wanted to quit halfway. With patience the last piece fit. Moral: Finish what you begin.",
    ),
    (
        "The Soft Voice Wins",
        "In an argument, shouting grew loud. One child spoke softly and clearly. Others calmed down and listened. Moral: Soft voices can stop big fights.",
    ),
    (
        "Watering Without Wasting",
        "A child watered plants with a small can, not a big hose. Plants got enough and water was saved. Moral: Use only what you need.",
    ),
    (
        "The First Medal",
        "After months of practice, a runner won a small medal. The child thanked the coach and teammates. Moral: Success is shared with helpers.",
    ),
    (
        "Saying Good Morning",
        "Every day a child said \"Good morning\" to the watchman and bus driver. Their faces lit up. Moral: Respect everyone you meet.",
    ),
    (
        "The Brave Hospital Visit",
        "A child feared the doctor. Holding Mother's hand, the child went bravely and got better soon. Moral: Facing fear with love makes you strong.",
    ),
    (
        "You Are Enough",
        "A little star thought it was too small. The moon said, \"Your light still guides travellers.\" The star shone proudly. Moral: You are enough, just as you are. Keep shining.",
    ),
]


def assert_hundred():
    assert len(STORIES) == 100, f"Expected 100 stories, got {len(STORIES)}"
    titles = [t for t, _ in STORIES]
    assert len(titles) == len(set(titles)), "Duplicate titles found"


def build_docx(path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(12)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("100 Motivational Stories for Class 2 Students")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = subtitle.add_run(
        "Short and simple stories for Grade 2 / 2nd Standard kids\n"
        "Read aloud at home or in class • One story a day"
    )
    sub.font.size = Pt(11)

    tips = doc.add_paragraph()
    tips.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tip_run = tips.add_run(
        "Tips: Read slowly • Ask \"What did we learn?\" • Praise kindness and effort"
    )
    tip_run.italic = True
    tip_run.font.size = Pt(10)

    doc.add_paragraph()

    for i, (story_title, body) in enumerate(STORIES, start=1):
        heading = doc.add_paragraph()
        h = heading.add_run(f"Story {i}: {story_title}")
        h.bold = True
        h.font.size = Pt(13)
        h.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

        para = doc.add_paragraph()
        p = para.add_run(body)
        p.font.size = Pt(12)
        doc.add_paragraph()

    doc.save(path)


def build_txt(path: Path) -> None:
    lines = [
        "100 Motivational Stories for Class 2 Students",
        "Short and simple stories for Grade 2 / 2nd Standard kids",
        "Read aloud at home or in class • One story a day",
        "",
        "Tips: Read slowly • Ask \"What did we learn?\" • Praise kindness and effort",
        "",
    ]
    for i, (title, body) in enumerate(STORIES, start=1):
        lines.append(f"Story {i}: {title}")
        lines.append(body)
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def build_html(path: Path) -> None:
    parts = [
        "<!DOCTYPE html>",
        '<html lang="en">',
        "<head>",
        '<meta charset="utf-8"/>',
        "<title>100 Motivational Stories for Class 2</title>",
        "<style>",
        "body{font-family:Georgia,serif;max-width:720px;margin:2rem auto;padding:0 1rem;line-height:1.55;color:#222}",
        "h1{color:#1f4e79;text-align:center}",
        ".sub,.tips{text-align:center;color:#444}",
        "h2{color:#2e75b6;font-size:1.15rem;margin-top:1.6rem}",
        "hr{border:none;border-top:1px solid #ddd;margin:2rem 0}",
        "</style>",
        "</head>",
        "<body>",
        "<h1>100 Motivational Stories for Class 2 Students</h1>",
        '<p class="sub">Short and simple stories for Grade 2 / 2nd Standard kids<br/>Read aloud at home or in class • One story a day</p>',
        '<p class="tips"><em>Tips: Read slowly • Ask “What did we learn?” • Praise kindness and effort</em></p>',
        "<hr/>",
    ]
    for i, (title, body) in enumerate(STORIES, start=1):
        parts.append(f"<h2>Story {i}: {title}</h2>")
        parts.append(f"<p>{body}</p>")
    parts.extend(["</body>", "</html>"])
    path.write_text("\n".join(parts), encoding="utf-8")


def build_rtf(path: Path) -> None:
    def esc(s: str) -> str:
        return (
            s.replace("\\", "\\\\")
            .replace("{", "\\{")
            .replace("}", "\\}")
            .replace("\n", "\\par ")
        )

    chunks = [
        r"{\rtf1\ansi\deff0",
        r"{\fonttbl{\f0 Calibri;}}",
        r"\f0\fs24",
        r"\qc\b\fs36 100 Motivational Stories for Class 2 Students\b0\fs24\par",
        r"\qc Short and simple stories for Grade 2 / 2nd Standard kids\par",
        r"\qc Read aloud at home or in class\line One story a day\par\par",
        r"\qc\i Tips: Read slowly \bullet Ask \"What did we learn?\" \bullet Praise kindness and effort\i0\par\par",
        r"\ql",
    ]
    for i, (title, body) in enumerate(STORIES, start=1):
        chunks.append(rf"\b Story {i}: {esc(title)}\b0\par")
        chunks.append(esc(body) + r"\par\par")
    chunks.append("}")
    path.write_text("".join(chunks), encoding="utf-8")


def main() -> None:
    assert_hundred()
    out_dir = Path(__file__).resolve().parent
    docx_path = out_dir / "100_Motivational_Stories_Class_2.docx"
    txt_path = out_dir / "100_Motivational_Stories_Class_2.txt"
    html_path = out_dir / "100_Motivational_Stories_Class_2.html"
    rtf_path = out_dir / "100_Motivational_Stories_Class_2.rtf"
    zip_path = out_dir / "100_Motivational_Stories_Class_2.zip"

    build_docx(docx_path)
    build_txt(txt_path)
    build_html(html_path)
    build_rtf(rtf_path)

    with ZipFile(zip_path, "w", ZIP_DEFLATED) as zf:
        for p in (docx_path, txt_path, html_path, rtf_path):
            zf.write(p, arcname=p.name)

    # Copy Word file to artifacts for easy download in Cursor Cloud
    artifacts = Path("/opt/cursor/artifacts")
    artifacts.mkdir(parents=True, exist_ok=True)
    artifact_docx = artifacts / docx_path.name
    artifact_zip = artifacts / zip_path.name
    artifact_docx.write_bytes(docx_path.read_bytes())
    artifact_zip.write_bytes(zip_path.read_bytes())

    print(f"Created {len(STORIES)} stories")
    print(f"DOCX: {docx_path} ({docx_path.stat().st_size} bytes)")
    print(f"ZIP:  {zip_path} ({zip_path.stat().st_size} bytes)")
    print(f"Artifact DOCX: {artifact_docx}")
    print(f"Artifact ZIP:  {artifact_zip}")


if __name__ == "__main__":
    main()
